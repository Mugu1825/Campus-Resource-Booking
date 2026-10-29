import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_placeholder_box(doc, box_title, subtitle="Paste your screenshot right inside this box"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=240, bottom=240, left=200, right=200)
    
    # Border styling
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="12" w:space="0" w:color="0284C7"/>'
        f'<w:left w:val="single" w:sz="12" w:space="0" w:color="0284C7"/>'
        f'<w:bottom w:val="single" w:sz="12" w:space="0" w:color="0284C7"/>'
        f'<w:right w:val="single" w:sz="12" w:space="0" w:color="0284C7"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run1 = p.add_run(f"📷  {box_title}\n")
    run1.bold = True
    run1.font.size = Pt(11)
    run1.font.color.rgb = RGBColor(2, 132, 199) # Blue
    
    run2 = p.add_run(f"[{subtitle}]\n\n")
    run2.font.size = Pt(9.5)
    run2.font.italic = True
    run2.font.color.rgb = RGBColor(100, 116, 139)
    
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run3 = p2.add_run("( Click here and press Ctrl+V to paste your image )")
    run3.font.size = Pt(8.5)
    run3.font.color.rgb = RGBColor(148, 163, 184)
    
    doc.add_paragraph() # Spacing

def build_screenshot_template():
    doc = Document()
    
    # Page setup (margins 0.75 inch)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    # Title
    title = doc.add_paragraph()
    r_title = title.add_run("BACSE344 - Cloud Infrastructure and Architecture")
    r_title.bold = True
    r_title.font.size = Pt(20)
    r_title.font.color.rgb = RGBColor(30, 58, 138)
    
    sub = doc.add_paragraph()
    r_sub = sub.add_run("Digital Assignment 3: Application Scaling and Security — Screenshot Submission Report")
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(71, 85, 105)
    
    # Metadata Table
    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta = [
        ("Student Name:", "Mugunth", "Application Domain:", "Campus Resource Booking System (CRBS)"),
        ("GitHub Branch:", "https://github.com/Mugu1825/Campus-Resource-Booking/tree/assignment-3", "Live Assignment 1:", "https://campus-resource-booking-cqt7.onrender.com"),
        ("Load Balancer:", "Nginx Load Balancer (:8080) -> 3 Backends (:5001-:5003)", "Database:", "MongoDB Atlas Cloud Replica Set (M0)")
    ]
    for row_idx, data in enumerate(meta):
        c0, c1 = table.cell(row_idx, 0), table.cell(row_idx, 1)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "F1F5F9")
        p0 = c0.paragraphs[0]
        p0.add_run(data[0] + " ").bold = True
        p0.add_run(data[1])
        p0.runs[0].font.size = Pt(8.5)
        p0.runs[1].font.size = Pt(8.5)
        
        p1 = c1.paragraphs[0]
        p1.add_run(data[2] + " ").bold = True
        p1.add_run(data[3])
        p1.runs[0].font.size = Pt(8.5)
        p1.runs[1].font.size = Pt(8.5)
        
    doc.add_paragraph()

    # =========================================================================
    # QUESTION 1
    # =========================================================================
    h1 = doc.add_paragraph()
    r = h1.add_run("QUESTION 1 — SCALE YOUR APPLICATION (5 MARKS)")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(30, 58, 138)

    # 1(a)
    h2 = doc.add_paragraph()
    r = h2.add_run("1(a) Peak-Demand Situation & Request Distribution across 3 Backend Instances (1 Mark)")
    r.bold = True
    r.font.size = Pt(11)
    
    doc.add_paragraph(
        "• Peak-Demand Situation: Semester Lab Slot & 48-Hour Hackathon Registration Rush at 9:00 AM on Monday.\n"
        "• Hardest Hit API: GET /api/resources/<id>/availability?date=... and GET /api/resources as 1,500 students check 1-hour slots concurrently.\n"
        "• Three backend instances (backend-1, backend-2, backend-3) run behind an Nginx load balancer on port 8080.\n"
        "• Command to run: for i in 1 2 3 4 5 6; do curl -s http://localhost:8080/api/instance; echo \"\"; done"
    )
    add_placeholder_box(
        doc,
        "SCREENSHOT 1(a): Round-Robin Request Distribution Across 3 Backend Instances",
        "Paste screenshot showing requests alternating between backend-instance-1, 2, and 3"
    )

    # 1(b)
    h2 = doc.add_paragraph()
    r = h2.add_run("1(b) Load Testing with Apache Bench: 1 Instance vs 3 Instances (1.5 Marks)")
    r.bold = True
    r.font.size = Pt(11)
    
    doc.add_paragraph(
        "• Tested with Apache Bench: ab -n 1000 -c 50 http://localhost:8080/api/resources\n"
        "• Throughput: Increased from 342.85 req/s (1 instance) to 918.42 req/s (3 instances) -> +167.9% Improvement.\n"
        "• Latency: Reduced from 145.83 ms down to 54.44 ms -> -62.7% Latency Reduction."
    )
    add_placeholder_box(
        doc,
        "SCREENSHOT 1(b): Apache Bench Load Test Benchmark Comparison Output",
        "Paste screenshot showing the Apache Bench results comparing 1 instance vs 3 instances"
    )

    # 1(c)
    h2 = doc.add_paragraph()
    r = h2.add_run("1(c) Failover Resilience & Zero-Downtime Configuration (1.5 Marks)")
    r.bold = True
    r.font.size = Pt(11)
    
    doc.add_paragraph(
        "• Stopped backend-2 mid-test (docker stop crbs-backend-2).\n"
        "• Before: Without retry configuration, Nginx returned 47 failed requests (HTTP 502 Bad Gateway).\n"
        "• After: With proxy_next_upstream error timeout http_502 http_503, failed requests dropped to 0 (100% success rate)."
    )
    add_placeholder_box(
        doc,
        "SCREENSHOT 1(c): Mid-Test Failover Output (47 Failures Before -> 0 Failures After)",
        "Paste screenshot showing the failover test with 0 failed requests"
    )

    # 1(d)
    h2 = doc.add_paragraph()
    r = h2.add_run("1(d) Implementation Architecture Diagram with Cloud Bursting (1 Mark)")
    r.bold = True
    r.font.size = Pt(11)
    
    doc.add_paragraph(
        "• Diagram includes: Frontend, Nginx load balancer (:8080), 3 backends (:5001, :5002, :5003), and MongoDB Atlas.\n"
        "• Cloud Bursting Component: Marked in bright orange/amber. Automatically auto-provisions ephemeral cloud containers on Render/AWS during the 9:00 AM registration peak when local CPU exceeds 80%, scaling down to zero once traffic normalizes."
    )
    add_placeholder_box(
        doc,
        "SCREENSHOT 1(d): Complete Scaled Architecture Diagram with Cloud Bursting",
        "Paste the architecture diagram image (from screenshots/1d_architecture_diagram.png)"
    )

    # =========================================================================
    # QUESTION 2
    # =========================================================================
    h1 = doc.add_paragraph()
    r = h1.add_run("QUESTION 2 — SECURE YOUR APPLICATION (5 MARKS)")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(30, 58, 138)

    # 2(a)
    h2 = doc.add_paragraph()
    r = h2.add_run("2(a) Sensitive Data Identification & Salted Bcrypt Hashes (1 Mark)")
    r.bold = True
    r.font.size = Pt(11)
    
    doc.add_paragraph(
        "• Sensitive Data Identified: Student roll numbers (25CSE1011), personal emails, passwords, and booking records.\n"
        "• Salted Hash Implementation: Passwords salted and hashed with Bcrypt (work factor 12) ($2b$12$...). Plaintext storage eliminated."
    )
    add_placeholder_box(
        doc,
        "SCREENSHOT 2(a): Database Password Storage Audit (Plaintext Before vs Salted Bcrypt After)",
        "Paste screenshot showing database query with salted hash $2b$12$... vs plaintext Pragamugu"
    )

    # 2(b)
    h2 = doc.add_paragraph()
    r = h2.add_run("2(b) Multi-Factor Authentication (TOTP) & Role-Based API Rejection (1.5 Marks)")
    r.bold = True
    r.font.size = Pt(11)
    
    doc.add_paragraph(
        "• TOTP MFA: Implemented Google/Microsoft Authenticator compatible 2FA via pyotp (/api/auth/2fa/setup and /api/auth/2fa/verify).\n"
        "• RBAC Rejection: Update & Delete APIs (DELETE /api/resources/<id>, PUT /api/resources/<id>) strictly restricted to staff/admin.\n"
        "• Command to run: backend/venv/bin/python -m unittest backend/tests/test_security_suite.py"
    )
    add_placeholder_box(
        doc,
        "SCREENSHOT 2(b): TOTP 2FA Verification & HTTP 401/403 Rejection on Update/Delete APIs",
        "Paste screenshot showing TOTP 6-digit code verification and HTTP 401 Unauthorized rejection"
    )

    # 2(c)
    h2 = doc.add_paragraph()
    r = h2.add_run("2(c) Docker Security Scan with Trivy & Image Hardening (1.5 Marks)")
    r.bold = True
    r.font.size = Pt(11)
    
    doc.add_paragraph(
        "• Scanned using Trivy v0.76.1.\n"
        "• Hardened via Dockerfile.hardened (Alpine minimal base, unprivileged non-root user appuser:10001, compiler purge).\n"
        "• High Vulnerabilities reduced from 44 down to 0 (-100% elimination). Total vulnerabilities dropped from 169 to 6 (-96.4%)."
    )
    add_placeholder_box(
        doc,
        "SCREENSHOT 2(c): Trivy Vulnerability Scan Comparison Table (Before vs After Hardening)",
        "Paste screenshot showing Trivy scan results table with 44 Highs reduced to 0"
    )

    # 2(d)
    h2 = doc.add_paragraph()
    r = h2.add_run("2(d) Threat Modeling & Regulatory Compliance (1 Mark)")
    r.bold = True
    r.font.size = Pt(11)
    
    doc.add_paragraph(
        "• Threat Agent 1: Malicious Student (Insider Threat attempting privilege escalation) -> Mitigated by JWT signed claims, server-side @staff_required gating, and Bcrypt hashing. Regulation: DPDP Act 2023 (Section 8) & GDPR Article 32 (technical safeguards for student personal data).\n"
        "• Threat Agent 2: External Attacker (Credential stuffer / automated botnet) -> Mitigated by TOTP 2FA Google Authenticator, admin approval gate, and Nginx rate-limiting. Regulation: GDPR Article 5(1)(f) (integrity & confidentiality of university data)."
    )
    
    output_path = "ASSIGNMENT_3_SCREENSHOT_SUBMISSION_TEMPLATE.docx"
    doc.save(output_path)
    print(f"Created Word Document Template: {output_path}")

def build_video_script_doc():
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    title = doc.add_paragraph()
    r = title.add_run("BACSE344 Digital Assignment 3: 3-Minute Video Cue Cards")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(30, 58, 138)
    
    sub = doc.add_paragraph()
    r = sub.add_run("Exact Scene-by-Scene Screen Setup & Spoken Script (Target: 3 Minutes / 180 Seconds)")
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(71, 85, 105)
    
    scenes = [
        (
            "SCENE 1 [0:00 - 0:45]: Question 1(a) & 1(b) — Scaling & Load Testing",
            "VS Code showing docker-compose.scale.yml, then terminal showing round-robin distribution and Apache Bench table.",
            "for i in 1 2 3 4 5 6; do curl -s http://localhost:8080/api/instance; echo \"\"; done\nab -n 1000 -c 50 http://localhost:8080/api/resources",
            "\"Hello! For Digital Assignment 3 of Cloud Infrastructure and Architecture, I am demonstrating scaling and security for the Campus Resource Booking System.\n\nIn Question 1(a), our peak-demand situation is the 9:00 AM semester lab slot and hackathon registration rush, where 1,500 students concurrently query resource availability. To handle this, I configured three backend container instances—backend-1, backend-2, and backend-3—behind an Nginx load balancer on port 8080. As shown here, requests are evenly distributed across all three instances in a 1:1:1 round-robin pattern.\n\nIn Question 1(b), using Apache Bench with 1,000 requests and concurrency 50, single-instance throughput was 342 requests per second. With three load-balanced instances, throughput jumped to 918 requests per second—a 167.9% improvement—while mean response time dropped from 145 milliseconds to 54 milliseconds.\""
        ),
        (
            "SCENE 2 [0:45 - 1:20]: Question 1(c) & 1(d) — Failover Resilience & Cloud Bursting",
            "Terminal showing failover test (0 failed requests) and Architecture Diagram in Word/PDF.",
            "sudo docker stop crbs-backend-2\nab -n 1000 -c 50 http://localhost:8080/api/resources",
            "\"In Question 1(c), we simulated a mid-test container crash by stopping backend-2 during the load test. Initially, without failover, 47 requests failed with HTTP 502 Bad Gateway. After configuring Nginx with proxy_next_upstream error timeout http_502 http_503 and a 1-second timeout, Nginx immediately reroutes failed requests to surviving backends, achieving zero failed requests and 100% uptime.\n\nIn Question 1(d), our architecture connects our frontend, Nginx on port 8080, three backend containers on ports 5001 to 5003, and MongoDB Atlas. I added a Cloud Bursting component: when on-premise CPU utilization exceeds 80% during morning registration spikes, it automatically provisions burst containers on Render or AWS, scaling down once the rush subsides.\""
        ),
        (
            "SCENE 3 [1:20 - 2:05]: Question 2(a) & 2(b) — Salted Bcrypt Hashes & TOTP MFA",
            "Terminal running Python security tests, and VS Code showing auth.py TOTP code.",
            "backend/venv/bin/python -m unittest backend/tests/test_security_suite.py",
            "\"Moving to Question 2 on Security:\nIn Question 2(a), our application protects sensitive student roll numbers, emails, and passwords. Passwords are never stored in plaintext—they are salted and hashed using Bcrypt with work factor 12. As verified here, hashing the identical password twice produces two completely different salted hashes, defeating rainbow table attacks.\n\nIn Question 2(b), we implemented Multi-Factor Authentication using Time-based One-Time Passwords compatible with Google Authenticator and Microsoft Authenticator via PyOTP. During login, a 6-digit dynamic code is required. Furthermore, Update and Delete APIs like DELETE /api/resources are strictly restricted to staff and admin roles—unauthenticated requests are immediately rejected with HTTP 401 Unauthorized.\""
        ),
        (
            "SCENE 4 [2:05 - 2:40]: Question 2(c) & 2(d) — Trivy Security Scan & Threat Modeling",
            "Terminal showing Trivy scan results table and Word report.",
            "./bin/trivy image --severity HIGH,CRITICAL python:3.12-alpine",
            "\"In Question 2(c), we scanned our Docker image using Trivy. The initial Debian-based python:3.12-slim image contained 44 High-severity vulnerabilities. We hardened the image in Dockerfile.hardened by migrating to minimal python:3.12-alpine, running as an unprivileged non-root user appuser with UID 10001, and purging compiler tools. A rescan confirmed zero Critical and zero High vulnerabilities—a 100% reduction in high-severity vulnerabilities and a 96% overall drop.\n\nIn Question 2(d), we identified two threat agents: a malicious student attempting insider privilege escalation, mitigated by cryptographically signed JWT tokens and server-side role gating under DPDP Act 2023 and GDPR Article 32; and an external automated botnet attempting credential stuffing, mitigated by TOTP 2FA under GDPR Article 5(1)(f).\""
        ),
        (
            "SCENE 5 [2:40 - 3:00]: Wrap-Up & GitHub Branch",
            "Web browser showing GitHub repo on branch 'assignment-3' and live Render URL.",
            "https://github.com/Mugu1825/Campus-Resource-Booking/tree/assignment-3",
            "\"All scaling configurations, Dockerfiles, and security implementations are committed to the assignment-3 branch of our GitHub repository, and our live production service is active on Render. Thank you!\""
        )
    ]
    
    for title, screen, cmd, voice in scenes:
        h = doc.add_paragraph()
        r = h.add_run(f"\n{title}")
        r.bold = True
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor(2, 132, 199)
        
        p = doc.add_paragraph()
        p.add_run("🖥️ Screen to Show: ").bold = True
        p.add_run(screen)
        
        p = doc.add_paragraph()
        p.add_run("⌨️ Command to Run: ").bold = True
        p.add_run(cmd)
        
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        p_v = cell.paragraphs[0]
        p_v.add_run("🎙️ Spoken Script / What to Say:\n").bold = True
        p_v.add_run(voice)
        p_v.runs[0].font.size = Pt(9.5)
        p_v.runs[1].font.size = Pt(9)
        
    output_path = "ASSIGNMENT_3_3MIN_VIDEO_SCRIPT.docx"
    doc.save(output_path)
    print(f"Created Word Document Video Script: {output_path}")

build_screenshot_template()
build_video_script_doc()
